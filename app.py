import re
import unicodedata
from io import BytesIO
from urllib.parse import urlencode

import pandas as pd
import streamlit as st
from openpyxl.styles import Font, PatternFill


st.set_page_config(
    page_title="Gerador de Links WhatsApp",
    page_icon="📲",
    layout="wide",
)


def normalizar_texto(texto):
    texto = unicodedata.normalize("NFKD", str(texto))
    texto = "".join(
        caractere for caractere in texto
        if not unicodedata.combining(caractere)
    )
    return re.sub(r"[^a-z0-9]", "", texto.lower())


def localizar_coluna(df, alternativas):
    mapa = {normalizar_texto(coluna): coluna for coluna in df.columns}
    for alternativa in alternativas:
        chave = normalizar_texto(alternativa)
        if chave in mapa:
            return mapa[chave]
    return None


def limpar_telefone(valor):
    if pd.isna(valor):
        return ""
    texto = str(valor).strip()
    if texto.endswith(".0"):
        texto = texto[:-2]
    return re.sub(r"\D", "", texto)


def gerar_link(telefone_aluno, mensagem):
    parametros = {
        "phone": telefone_aluno,
        "text": mensagem,
        "type": "phone_number",
        "app_absent": "0",
    }
    return "https://api.whatsapp.com/send/?" + urlencode(parametros)


def criar_excel(df):
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Alunos")
        planilha = writer.sheets["Alunos"]

        for celula in planilha[1]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill(
                fill_type="solid",
                fgColor="167D46",
            )

        nomes_colunas = list(df.columns)
        if "Link WhatsApp" in nomes_colunas:
            indice_link = nomes_colunas.index("Link WhatsApp") + 1
            for linha in range(2, planilha.max_row + 1):
                celula = planilha.cell(row=linha, column=indice_link)
                if celula.value:
                    celula.hyperlink = celula.value
                    celula.style = "Hyperlink"

        planilha.freeze_panes = "A2"
        planilha.auto_filter.ref = planilha.dimensions

        for coluna in planilha.columns:
            letra = coluna[0].column_letter
            maior = max(
                len(str(celula.value or ""))
                for celula in coluna[:30]
            )
            planilha.column_dimensions[letra].width = min(
                max(maior + 2, 15), 40
            )

        if "Link WhatsApp" in nomes_colunas:
            letra_link = planilha.cell(
                row=1,
                column=nomes_colunas.index("Link WhatsApp") + 1,
            ).column_letter
            planilha.column_dimensions[letra_link].width = 25

    buffer.seek(0)
    return buffer.getvalue()


st.title("📲 Gerador de Links WhatsApp")
st.write(
    "Envie a planilha de alunos e gere um Excel com links clicáveis "
    "para mensagens personalizadas no WhatsApp."
)
st.info("A planilha precisa conter colunas equivalentes a Nome, Curso e Telefone.")

with st.expander("⚙️ Configurações do instrutor", expanded=True):
    col1, col2, col3 = st.columns(3)
    with col1:
        nome_instrutor = st.text_input(
            "Nome ou identificação do instrutor",
            value="Instrutor",
        )
    with col2:
        telefone_instrutor = st.text_input(
            "WhatsApp do instrutor (com país e DDD)",
            value="5562999999999",
            help="Informe somente números, incluindo código do país e DDD.",
        )
    with col3:
        instituicao = st.text_input(
            "Nome da instituição",
            value="Grau Técnico",
        )

st.subheader("💬 Tipo de mensagem")
tipo_mensagem = st.radio(
    "Escolha o público da mensagem:",
    [
        "1: Mensagem customizada para alunos com presença",
        "2: Mensagem customizada para alunos faltosos",
    ],
    horizontal=False,
)

mensagem_presenca = (
    "Olá, {nome}! Tudo bem? 😊\n\n"
    "Aqui é {nome_instrutor}, novo instrutor do Curso Interativo da "
    "{instituicao}.\n\n"
    "Estou entrando em contato para acompanhar sua participação no curso "
    "{curso} e reforçar a importância de manter a frequência e continuar "
    "avançando nos estudos.\n\n"
    "Este número será nosso novo canal de contato para comunicados sobre "
    "a rotina do Curso Interativo.\n\n"
    "Conte comigo!\n\n"
    "Atenciosamente,\n"
    "{nome_instrutor}\n"
    "Instrutor do Curso Interativo — {instituicao}"
)

mensagem_falta = (
    "Olá, {nome}! Tudo bem? Aqui é {nome_instrutor}, instrutor do Curso "
    "Interativo da {instituicao}.\n\n"
    "Percebemos sua ausência nas aulas do curso {curso} e gostaríamos de "
    "saber se está tudo bem.\n\n"
    "Sua participação é importante para dar continuidade aos estudos. "
    "Quando puder, entre em contato com a equipe da unidade para verificar "
    "as orientações necessárias para sua rotina de aulas.\n\n"
    "Estamos à disposição!\n\n"
    "Atenciosamente,\n"
    "{nome_instrutor}\n"
    "Instrutor do Curso Interativo — {instituicao}"
)

modelo_padrao = (
    mensagem_presenca
    if tipo_mensagem.startswith("1:")
    else mensagem_falta
)
mensagem_modelo = st.text_area(
    "Mensagem (você pode editar antes de gerar os links)",
    value=modelo_padrao,
    height=260,
    help=(
        "Campos automáticos disponíveis: {nome}, {curso}, "
        "{nome_instrutor} e {instituicao}."
    ),
)
st.caption(
    "Os campos entre chaves serão preenchidos automaticamente com os dados "
    "de cada aluno e as configurações do instrutor."
)

arquivo = st.file_uploader(
    "📄 Selecione a planilha de alunos",
    type=["xlsx", "csv"],
    help="Formatos aceitos: Excel (.xlsx) ou CSV (.csv).",
)

if arquivo is not None:
    try:
        if arquivo.name.lower().endswith(".csv"):
            df = pd.read_csv(
                arquivo,
                dtype=str,
                sep=None,
                engine="python",
                encoding="utf-8-sig",
            )
        else:
            df = pd.read_excel(arquivo, dtype=str, engine="openpyxl")

        df.columns = [str(coluna).strip() for coluna in df.columns]

        if df.empty:
            st.warning("A planilha está vazia.")
        else:
            col_nome = localizar_coluna(
                df, ["Nome", "Nome do aluno", "Aluno"]
            )
            col_curso = localizar_coluna(
                df, ["Curso", "Nome do curso"]
            )
            col_telefone = localizar_coluna(
                df,
                [
                    "Telefone", "WhatsApp", "Celular",
                    "Telefone do aluno", "Número", "Numero",
                ],
            )

            if not all([col_nome, col_curso, col_telefone]):
                st.error(
                    "Não consegui identificar todas as colunas necessárias."
                )
                st.write("Colunas encontradas:", list(df.columns))
                st.write(
                    "São necessárias colunas equivalentes a Nome, Curso e Telefone."
                )
            else:
                st.success(f"Planilha carregada: {len(df)} registros.")
                st.subheader("Prévia dos alunos")
                st.dataframe(
                    df[[col_nome, col_curso, col_telefone]].head(10),
                    use_container_width=True,
                    hide_index=True,
                )

                if st.button(
                    "🔗 Gerar links WhatsApp",
                    type="primary",
                    use_container_width=True,
                ):
                    telefone_padrao = limpar_telefone(telefone_instrutor)

                    if not nome_instrutor.strip():
                        st.error("Informe o nome do instrutor.")
                    elif not instituicao.strip():
                        st.error("Informe o nome da instituição.")
                    elif not (10 <= len(telefone_padrao) <= 15):
                        st.error(
                            "O telefone do instrutor é inválido. "
                            "Informe país, DDD e número."
                        )
                    else:
                        resultado = df.copy()
                        links = []
                        status = []

                        for _, aluno in resultado.iterrows():
                            nome = (
                                ""
                                if pd.isna(aluno[col_nome])
                                else str(aluno[col_nome]).strip()
                            )
                            curso = (
                                ""
                                if pd.isna(aluno[col_curso])
                                else str(aluno[col_curso]).strip()
                            )
                            telefone_aluno = limpar_telefone(aluno[col_telefone])

                            if not nome or not curso:
                                links.append("")
                                status.append("Dados incompletos")
                                continue

                            if not (10 <= len(telefone_aluno) <= 15):
                                links.append("")
                                status.append("Telefone inválido")
                                continue

                            try:
                                mensagem = mensagem_modelo.format(
                                    nome=nome,
                                    curso=curso,
                                    nome_instrutor=nome_instrutor.strip(),
                                    instituicao=instituicao.strip(),
                                )
                            except KeyError as erro:
                                links.append("")
                                status.append(
                                    f"Campo desconhecido no modelo: {erro}"
                                )
                                continue

                            links.append(gerar_link(telefone_aluno, mensagem))
                            status.append("OK")

                        resultado["Tipo de mensagem"] = (
                            "Alunos com presença"
                            if tipo_mensagem.startswith("1:")
                            else "Alunos faltosos"
                        )
                        resultado["Link WhatsApp"] = links
                        resultado["Status"] = status
                        st.session_state["resultado"] = resultado
                        st.session_state["tipo_resultado"] = (
                            "presenca"
                            if tipo_mensagem.startswith("1:")
                            else "faltosos"
                        )

    except Exception as erro:
        st.error(f"Não foi possível processar a planilha: {erro}")


if "resultado" in st.session_state:
    resultado = st.session_state["resultado"]
    total = len(resultado)
    validos = int((resultado["Status"] == "OK").sum())
    invalidos = total - validos

    st.divider()
    st.header("📊 Resultado do processamento")
    st.caption(
        "Tipo selecionado: "
        + (
            "Mensagem para alunos com presença"
            if st.session_state.get("tipo_resultado") == "presenca"
            else "Mensagem para alunos faltosos"
        )
    )

    coluna1, coluna2, coluna3 = st.columns(3)
    coluna1.metric("Total de alunos", total)
    coluna2.metric("Links gerados", validos)
    coluna3.metric("Registros para revisar", invalidos)

    st.dataframe(
        resultado,
        use_container_width=True,
        hide_index=True,
        column_config={
            "Link WhatsApp": st.column_config.LinkColumn(
                "Link WhatsApp",
                display_text="Abrir conversa",
            )
        },
    )

    arquivo_excel = criar_excel(resultado)
    st.download_button(
        label="⬇️ Baixar planilha com os links",
        data=arquivo_excel,
        file_name=(
            "alunos_com_presenca_links.xlsx"
            if st.session_state.get("tipo_resultado") == "presenca"
            else "alunos_faltosos_links.xlsx"
        ),
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        type="primary",
        use_container_width=True,
    )
    st.caption(
        "Os links abrem a conversa do WhatsApp com a mensagem preenchida. "
        "É necessário confirmar o envio no WhatsApp."
    )

st.divider()
st.caption("Gerador de Links WhatsApp — Grau Técnico")
